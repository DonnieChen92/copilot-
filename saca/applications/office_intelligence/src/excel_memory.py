import json
from typing import List, Dict, Any
import openpyxl
from openpyxl import Workbook
import os

class ExcelMemory:
    """
    Excel-based Memory and Knowledge Graph Storage.
    基于 Excel 的记忆与知识图谱存储。
    基於 Excel 的記憶與知識圖譜存儲。
    """

    def __init__(self, filepath: str = "memory.xlsx"):
        """
        Initialize the Excel Memory system.
        初始化 Excel 记忆系统。
        初始化 Excel 記憶系統。
        """
        self.filepath = filepath
        if not os.path.exists(filepath):
            self.wb = Workbook()
            self.wb.active.title = "Index"
            self.wb.create_sheet("Embeddings") # For vector storage
            self.wb.create_sheet("KnowledgeGraph") # For KG nodes/edges
            self.save()
        else:
            self.wb = openpyxl.load_workbook(filepath)

    def save(self):
        """Save the workbook / 保存工作簿 / 保存工作簿"""
        self.wb.save(self.filepath)

    def store_embedding(self, text: str, vector: List[float]):
        """
        Store text and its vector embedding.
        存储文本及其向量嵌入。
        存儲文本及其向量嵌入。

        :param text: The source text.
        :param vector: The embedding vector.
        """
        sheet = self.wb["Embeddings"]
        # Convert vector to JSON string for storage
        # 将向量转换为 JSON 字符串以进行存储
        # 將向量轉換為 JSON 字串以進行存儲
        vector_str = json.dumps(vector)
        sheet.append([text, vector_str])
        self.save()
        print(f"Stored embedding for: {text[:20]}...")

    def retrieve_similar(self, query_vector: List[float], top_k: int = 1) -> List[str]:
        """
        Retrieve text with similar embeddings (Mock implementation).
        检索具有相似嵌入的文本（模拟实现）。
        檢索具有相似嵌入的文本（模擬實現）。
        """
        # In a real implementation, we would compute cosine similarity here.
        # 在实际实现中，我们将在此处计算余弦相似度。
        # 在實際實現中，我們將在此處計算餘弦相似度。
        sheet = self.wb["Embeddings"]
        results = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0]:
                results.append(row[0])
        return results[:top_k]

    def add_knowledge_node(self, node_id: str, label: str, properties: Dict[str, Any]):
        """
        Add a node to the Knowledge Graph sheet.
        向知识图谱工作表添加节点。
        向知識圖譜工作表添加節點。
        """
        sheet = self.wb["KnowledgeGraph"]
        props_str = json.dumps(properties, ensure_ascii=False)
        sheet.append([node_id, label, props_str])
        self.save()

if __name__ == "__main__":
    mem = ExcelMemory("saca_test_memory.xlsx")
    mem.store_embedding("Project Plan V1", [0.1, 0.5, 0.9])
    mem.add_knowledge_node("N001", "Person", {"name": "Donnie", "role": "Architect"})
